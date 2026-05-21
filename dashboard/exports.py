from datetime import date
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

_COLOR_HEADER = HexColor('#1a1a2e')
_COLOR_ALT = HexColor('#f8f9fa')
_COLOR_ROJO_BG = HexColor('#ffebee')

_ESTADO_LABELS = {
    'pendiente': 'Pendiente',
    'en_preparacion': 'En preparación',
    'listo': 'Listo',
    'entregado': 'Entregado',
    'cancelado': 'Cancelado',
}


def _base_table_style():
    return [
        ('BACKGROUND', (0, 0), (-1, 0), _COLOR_HEADER),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('TOPPADDING', (0, 0), (-1, 0), 8),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
        ('TOPPADDING', (0, 1), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, HexColor('#dee2e6')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]


def _pdf_doc(buffer):
    return SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=1.5 * cm,
        leftMargin=1.5 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )


def _header_styles(styles):
    title = ParagraphStyle(
        'rapp_title',
        parent=styles['Heading1'],
        fontSize=18,
        textColor=_COLOR_HEADER,
        spaceAfter=4,
    )
    subtitle = ParagraphStyle(
        'rapp_subtitle',
        parent=styles['Normal'],
        fontSize=11,
        textColor=HexColor('#555555'),
        spaceAfter=2,
    )
    footer = ParagraphStyle(
        'rapp_footer',
        parent=styles['Normal'],
        fontSize=8,
        textColor=HexColor('#999999'),
    )
    return title, subtitle, footer


def generar_pdf_pedidos(pedidos, fecha_desde, fecha_hasta):
    """Genera PDF con tabla de pedidos y fila de total. Retorna BytesIO."""
    buffer = BytesIO()
    doc = _pdf_doc(buffer)
    styles = getSampleStyleSheet()
    title_s, subtitle_s, footer_s = _header_styles(styles)
    elements = []

    elements.append(Paragraph('RestaurApp', title_s))
    elements.append(Paragraph('Reporte de Pedidos', subtitle_s))
    elements.append(Paragraph(
        f'Período: {fecha_desde.strftime("%d/%m/%Y")} — {fecha_hasta.strftime("%d/%m/%Y")}',
        subtitle_s,
    ))
    elements.append(Spacer(1, 0.5 * cm))

    headers = ['N°', 'Cliente', 'Fecha', 'Platos', 'Total', 'Estado']
    col_widths = [1.2 * cm, 3.5 * cm, 2.8 * cm, 6.5 * cm, 2.2 * cm, 2.5 * cm]
    rows = [headers]
    total_general = 0

    for pedido in pedidos:
        platos_str = ', '.join(
            f'{d.plato.nombre} x{d.cantidad}' for d in pedido.detalles.all()
        ) or '—'
        rows.append([
            str(pedido.pk),
            pedido.cliente.get_full_name() or pedido.cliente.username,
            pedido.fecha_creacion.strftime('%d/%m/%Y %H:%M'),
            platos_str,
            f'${pedido.total}',
            _ESTADO_LABELS.get(pedido.estado, pedido.estado),
        ])
        total_general += pedido.total

    rows.append(['', '', '', 'TOTAL GENERAL', f'${total_general:.2f}', ''])

    table = Table(rows, colWidths=col_widths, repeatRows=1)
    style = _base_table_style()

    for i in range(1, len(rows) - 1):
        if i % 2 == 0:
            style.append(('BACKGROUND', (0, i), (-1, i), _COLOR_ALT))

    last = len(rows) - 1
    style.extend([
        ('BACKGROUND', (0, last), (-1, last), _COLOR_HEADER),
        ('TEXTCOLOR', (0, last), (-1, last), colors.white),
        ('FONTNAME', (0, last), (-1, last), 'Helvetica-Bold'),
    ])

    table.setStyle(TableStyle(style))
    elements.append(table)
    elements.append(Spacer(1, 0.5 * cm))
    elements.append(Paragraph(
        f'Generado el {date.today().strftime("%d/%m/%Y")} — RestaurApp',
        footer_s,
    ))

    doc.build(elements)
    buffer.seek(0)
    return buffer


def generar_excel_pedidos(pedidos, fecha_desde, fecha_hasta):
    """Genera Excel con hojas Pedidos y Resumen. Retorna BytesIO."""
    wb = openpyxl.Workbook()
    header_fill = PatternFill(start_color='1A1A2E', end_color='1A1A2E', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)
    alt_fill = PatternFill(start_color='F8F9FA', end_color='F8F9FA', fill_type='solid')
    center = Alignment(horizontal='center', vertical='center')

    # ── Hoja Pedidos ──────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = 'Pedidos'

    headers = ['N° Pedido', 'Cliente', 'Mesero', 'Fecha', 'Total', 'Estado', 'Método de pago']
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center

    totales_por_estado: dict[str, float] = {}
    total_general = 0.0

    for pedido in pedidos:
        mesero = ''
        if pedido.mesero:
            mesero = pedido.mesero.get_full_name() or pedido.mesero.username
        estado_label = _ESTADO_LABELS.get(pedido.estado, pedido.estado)
        ws.append([
            pedido.pk,
            pedido.cliente.get_full_name() or pedido.cliente.username,
            mesero,
            pedido.fecha_creacion.strftime('%d/%m/%Y %H:%M'),
            float(pedido.total),
            estado_label,
            pedido.get_metodo_pago_display(),
        ])
        totales_por_estado[estado_label] = (
            totales_por_estado.get(estado_label, 0.0) + float(pedido.total)
        )
        total_general += float(pedido.total)

    for i, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row), start=2):
        if i % 2 == 0:
            for cell in row:
                cell.fill = alt_fill

    for col_idx, col in enumerate(ws.columns, 1):
        max_len = max((len(str(c.value)) if c.value else 0 for c in col), default=0)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 40)

    # ── Hoja Resumen ──────────────────────────────────────────────────────────
    ws2 = wb.create_sheet('Resumen')
    ws2.append(['Período', f'{fecha_desde.strftime("%d/%m/%Y")} — {fecha_hasta.strftime("%d/%m/%Y")}'])
    ws2.append([])
    ws2.append(['Estado', 'Total ($)'])
    for cell in ws2[3]:
        cell.fill = header_fill
        cell.font = header_font

    for estado, total in totales_por_estado.items():
        ws2.append([estado, round(total, 2)])

    ws2.append([])
    ws2.append(['TOTAL GENERAL', round(total_general, 2)])
    for cell in ws2[ws2.max_row]:
        cell.font = Font(bold=True)

    ws2.column_dimensions['A'].width = 22
    ws2.column_dimensions['B'].width = 15

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generar_pdf_inventario(ingredientes):
    """Genera PDF del inventario, con filas en rojo para stock bajo. Retorna BytesIO."""
    buffer = BytesIO()
    doc = _pdf_doc(buffer)
    styles = getSampleStyleSheet()
    title_s, subtitle_s, footer_s = _header_styles(styles)
    elements = []

    elements.append(Paragraph('RestaurApp', title_s))
    elements.append(Paragraph('Estado del Inventario', subtitle_s))
    elements.append(Paragraph(
        f'Generado el {date.today().strftime("%d/%m/%Y")}',
        subtitle_s,
    ))
    elements.append(Spacer(1, 0.5 * cm))

    headers = ['Ingrediente', 'Unidad', 'Stock actual', 'Stock mínimo', 'Estado']
    col_widths = [7 * cm, 3 * cm, 3 * cm, 3 * cm, 2.5 * cm]
    rows = [headers]
    stock_bajo_rows = []

    for i, ing in enumerate(ingredientes, start=1):
        bajo = ing.stock_actual <= ing.stock_minimo
        rows.append([
            ing.nombre,
            ing.unidad_medida,
            str(ing.stock_actual),
            str(ing.stock_minimo),
            'BAJO' if bajo else 'OK',
        ])
        if bajo:
            stock_bajo_rows.append(i)

    table = Table(rows, colWidths=col_widths, repeatRows=1)
    style = _base_table_style()

    for i in range(1, len(rows)):
        if i % 2 == 0 and i not in stock_bajo_rows:
            style.append(('BACKGROUND', (0, i), (-1, i), _COLOR_ALT))

    for i in stock_bajo_rows:
        style.append(('BACKGROUND', (0, i), (-1, i), _COLOR_ROJO_BG))
        style.append(('TEXTCOLOR', (4, i), (4, i), HexColor('#c62828')))
        style.append(('FONTNAME', (4, i), (4, i), 'Helvetica-Bold'))

    table.setStyle(TableStyle(style))
    elements.append(table)
    elements.append(Spacer(1, 0.5 * cm))
    elements.append(Paragraph(
        f'Generado el {date.today().strftime("%d/%m/%Y")} — RestaurApp',
        footer_s,
    ))

    doc.build(elements)
    buffer.seek(0)
    return buffer


def generar_excel_inventario(ingredientes, movimientos):
    """Genera Excel con hojas Stock actual y Movimientos. Retorna BytesIO."""
    wb = openpyxl.Workbook()
    header_fill = PatternFill(start_color='1A1A2E', end_color='1A1A2E', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)
    rojo_fill = PatternFill(start_color='FFEBEE', end_color='FFEBEE', fill_type='solid')
    rojo_font = Font(color='C62828', bold=True)
    alt_fill = PatternFill(start_color='F8F9FA', end_color='F8F9FA', fill_type='solid')

    # ── Hoja Stock actual ─────────────────────────────────────────────────────
    ws = wb.active
    ws.title = 'Stock actual'
    ws.append(['Ingrediente', 'Unidad de medida', 'Stock actual', 'Stock mínimo', 'Estado'])
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    for i, ing in enumerate(ingredientes, start=2):
        bajo = ing.stock_actual <= ing.stock_minimo
        ws.append([
            ing.nombre,
            ing.unidad_medida,
            float(ing.stock_actual),
            float(ing.stock_minimo),
            'BAJO' if bajo else 'OK',
        ])
        if bajo:
            for cell in ws[i]:
                cell.fill = rojo_fill
            ws.cell(row=i, column=5).font = rojo_font
        elif i % 2 == 0:
            for cell in ws[i]:
                cell.fill = alt_fill

    for col_idx, col in enumerate(ws.columns, 1):
        max_len = max((len(str(c.value)) if c.value else 0 for c in col), default=0)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 35)

    # ── Hoja Movimientos ──────────────────────────────────────────────────────
    ws2 = wb.create_sheet('Movimientos')
    ws2.append(['Fecha', 'Ingrediente', 'Tipo', 'Cantidad', 'Stock resultante', 'Responsable', 'Motivo'])
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font

    for i, mov in enumerate(movimientos, start=2):
        ws2.append([
            mov.fecha.strftime('%d/%m/%Y %H:%M'),
            mov.ingrediente.nombre,
            mov.get_tipo_display(),
            float(mov.cantidad),
            float(mov.stock_resultante),
            mov.responsable.get_full_name() or mov.responsable.username,
            mov.motivo,
        ])
        if i % 2 == 0:
            for cell in ws2[i]:
                cell.fill = alt_fill

    for col_idx, col in enumerate(ws2.columns, 1):
        max_len = max((len(str(c.value)) if c.value else 0 for c in col), default=0)
        ws2.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 35)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
